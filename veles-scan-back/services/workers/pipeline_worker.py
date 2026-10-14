from __future__ import annotations

import logging
import os
import json
import re
import select
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from uuid import UUID
from collections.abc import Callable

from prometheus_client import start_http_server
from services.api.jobs_repo import append_debug_log, ensure_schema, get_job, get_job_cancel_reason, list_worker_heartbeats, update_job_state
from services.shared.logging_utils import setup_logging
from services.shared.metrics import jobs_processed_total, stage_duration_seconds
from services.shared.models import JobStatus
from services.shared.rmq import consume_json, publish_json
from services.shared.settings import SETTINGS
from services.shared.storage import download_file, ensure_bucket, list_keys, upload_file
from services.shared.worker_heartbeat import WorkerHeartbeatPublisher


setup_logging()
log = logging.getLogger("pipeline_worker")
ROOT = Path(__file__).resolve().parents[2]
OKVED_RE = re.compile(r"^\d{2}(?:\.\d+)*$")
STAGE03_SUCCESS_RE = re.compile(r"Успешно обработано:\s*(.+)$")
STAGE03_PROCESSING_RE = re.compile(r"Обработка компании:\s*(.+)$")


class JobCancelledError(RuntimeError):
    pass


def _normalize_okved_codes(raw_codes: object) -> list[str]:
    if not isinstance(raw_codes, list):
        return []
    result: list[str] = []
    seen: set[str] = set()
    for item in raw_codes:
        code = str(item).strip()
        if not OKVED_RE.fullmatch(code) or code in seen:
            continue
        seen.add(code)
        result.append(code)
    return result


def _emit(job_id: str, status: JobStatus, stage: str, progress: int, message: str, payload: dict | None = None) -> None:
    publish_json(
        SETTINGS.queue_status,
        {
            "job_id": job_id,
            "status": status.value,
            "stage": stage,
            "progress": progress,
            "message": message,
            "payload": payload or {},
        },
    )


def _append_job_log(
    job_uuid: UUID,
    *,
    level: str,
    message: str,
    payload: dict | None = None,
) -> None:
    append_debug_log(
        job_id=job_uuid,
        service="pipeline_worker",
        level=level,
        message=message,
        payload=payload,
    )


def _raise_if_cancel_requested(job_uuid: UUID) -> None:
    cancel_reason = get_job_cancel_reason(job_uuid)
    if cancel_reason:
        raise JobCancelledError(cancel_reason)


def _terminate_process(proc: subprocess.Popen[str], timeout_sec: float = 10.0) -> None:
    if proc.poll() is not None:
        return
    proc.terminate()
    try:
        proc.wait(timeout=timeout_sec)
    except subprocess.TimeoutExpired:
        proc.kill()
        proc.wait(timeout=5.0)


def _count_stage03_total_sites(base_dir: Path) -> int | None:
    workbook_path = base_dir / "active_companies_for_gpt.xlsx"
    if not workbook_path.exists():
        return None
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        worksheet = workbook.active
        total = max((worksheet.max_row or 1) - 1, 0)
        workbook.close()
        return total
    except Exception as exc:
        log.warning("Failed to count total sites from %s: %s", workbook_path, exc)
        return None


def _update_stage03_runtime_progress(
    job_uuid: UUID,
    *,
    processed_sites: int,
    total_sites: int,
    current_site_url: str | None,
    stage_started_at: float,
) -> None:
    if total_sites <= 0:
        return

    processed = max(0, min(processed_sites, total_sites))
    elapsed = max(time.time() - stage_started_at, 0.001)
    rate = processed / elapsed if processed > 0 else 0.0
    remaining = max(total_sites - processed, 0)
    eta_seconds = int(round(remaining / rate)) if rate > 0 and remaining > 0 else 0

    stage_min = 30
    stage_max = 44
    stage_span = max(stage_max - stage_min, 1)
    progress_percent = stage_min + int((processed / total_sites) * stage_span)
    progress_percent = max(stage_min, min(stage_max, progress_percent))

    update_job_state(
        job_uuid,
        progress=progress_percent,
        metadata_patch={
            "runtime_progress": {
                "processed_sites": processed,
                "total_sites": total_sites,
                "eta_seconds": eta_seconds,
                "current_site_url": current_site_url,
                "updated_at": datetime.now(timezone.utc).isoformat(),
            }
        },
    )


def _run_stage(
    job_uuid: UUID,
    script_name: str,
    env: dict[str, str],
    *,
    on_line: Callable[[str], None] | None = None,
) -> None:
    started = time.perf_counter()
    cmd = [sys.executable, str(ROOT / script_name)]
    proc = subprocess.Popen(
        cmd,
        cwd=str(ROOT),
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
    )
    assert proc.stdout is not None
    recent_lines: list[str] = []
    try:
        while True:
            _raise_if_cancel_requested(job_uuid)
            if proc.poll() is not None:
                break

            ready, _, _ = select.select([proc.stdout], [], [], 1.0)
            if not ready:
                continue

            line = proc.stdout.readline()
            if not line:
                continue
            chunk = line.rstrip()
            if not chunk:
                continue
            recent_lines.append(chunk)
            if len(recent_lines) > 50:
                recent_lines.pop(0)
            log.info("%s | %s", script_name, chunk)
            _append_job_log(job_uuid, level="DEBUG", message=chunk, payload={"stage_script": script_name})
            if on_line:
                on_line(chunk)
    except JobCancelledError:
        _terminate_process(proc)
        raise

    ret = proc.wait()
    if ret != 0:
        # Bubble up the most relevant stage output line to make user-facing errors actionable.
        stage_hint = ""
        for candidate in reversed(recent_lines):
            if "No rows left after applying intent filters" in candidate:
                stage_hint = candidate
                break
        if not stage_hint and recent_lines:
            stage_hint = recent_lines[-1]
        if stage_hint:
            raise RuntimeError(f"{script_name} failed: {stage_hint}")
        raise subprocess.CalledProcessError(ret, cmd)
    stage_duration_seconds.labels(stage=script_name).observe(time.perf_counter() - started)


def _process(payload: dict) -> None:
    job_id = payload["job_id"]
    job_uuid = UUID(job_id)
    existing = get_job(job_uuid)
    if existing and existing.status in {JobStatus.FAILED, JobStatus.COMPLETED}:
        # Drop stale queue messages for jobs that already reached terminal state.
        log.warning("Skipping stale pipeline message for terminal job %s (%s)", job_id, existing.status.value)
        return
    if existing and existing.status == JobStatus.RUNNING:
        heartbeats = list_worker_heartbeats(limit=500)
        already_processing_by_other = any(
            hb.get("worker_role") == "pipeline_worker"
            and hb.get("status") == "busy"
            and str(hb.get("current_job_id") or "") == job_id
            and hb.get("worker_id") != HEARTBEAT.worker_id
            for hb in heartbeats
        )
        if already_processing_by_other:
            log.warning("Skipping duplicate pipeline message for running job %s (another worker is already busy)", job_id)
            _append_job_log(
                job_uuid,
                level="WARNING",
                message="Skipped duplicate queue message: job is already being processed by another worker",
            )
            return
    input_bucket = payload["input_bucket"]
    input_key = payload["input_key"]
    job_type = payload.get("job_type", "enrich_only")
    metadata = payload.get("metadata", {}) or {}
    build_from_checko = bool(metadata.get("build_from_checko", False))
    selected_okved = _normalize_okved_codes(metadata.get("okved_codes"))

    base = Path(SETTINGS.pipeline_artifacts_root) / "jobs" / job_id
    base.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env["PIPELINE_ARTIFACTS_DIR"] = str(base)
    output_fields = metadata.get("output_fields")
    if isinstance(output_fields, list):
        env["PIPELINE_OUTPUT_FIELDS_JSON"] = json.dumps(output_fields, ensure_ascii=False)
    data_filters = metadata.get("data_filters")
    if isinstance(data_filters, list):
        env["PIPELINE_INTENT_FILTERS_JSON"] = json.dumps(data_filters, ensure_ascii=False)
    ignored_conditions = metadata.get("ignored_intent_conditions")
    if isinstance(ignored_conditions, list):
        env["PIPELINE_IGNORED_CONDITIONS_JSON"] = json.dumps(ignored_conditions, ensure_ascii=False)

    try:
        _raise_if_cancel_requested(job_uuid)
        HEARTBEAT.set_busy(job_uuid, "Pipeline job in progress")
        _append_job_log(
            job_uuid,
            level="INFO",
            message="Pipeline worker started processing job",
            payload={
                "job_type": job_type,
                "build_from_checko": build_from_checko,
                "okved_codes": selected_okved,
                "data_filters": data_filters if isinstance(data_filters, list) else [],
                "ignored_intent_conditions": ignored_conditions if isinstance(ignored_conditions, list) else [],
            },
        )
        if selected_okved:
            log.info("Using OKVED codes for job %s: %s", job_id, ", ".join(selected_okved))
            _append_job_log(
                job_uuid,
                level="INFO",
                message=f"Using OKVED codes: {', '.join(selected_okved)}",
                payload={"okved_codes_count": len(selected_okved)},
            )
        update_job_state(job_uuid, status=JobStatus.RUNNING, progress=1)
        _emit(job_id, JobStatus.RUNNING, "start", 1, "Pipeline started")

        if not build_from_checko:
            if not input_bucket or not input_key:
                raise ValueError("input_bucket and input_key are required for non-checko jobs")
            input_local = base / "uploaded_input.xlsx"
            download_file(input_bucket, input_key, input_local)
            target_input = base / "active_companies_for_gpt.xlsx"
            target_input.write_bytes(input_local.read_bytes())

        run_stage_02 = build_from_checko or (
            job_type == "checko_refresh_and_enrich" and not metadata.get("skip_stage_02", False)
        )
        if run_stage_02:
            # Sync latest checko snapshots from MinIO to local artifacts/checko_raw
            latest_keys = list_keys(SETTINGS.minio_bucket_checko, "checko/latest/")
            checko_raw_dir = base / "checko_raw"
            checko_raw_dir.mkdir(parents=True, exist_ok=True)
            selected_set = set(selected_okved)
            downloaded = 0
            for key in latest_keys:
                if not key.endswith(".csv"):
                    continue
                local_name = key.split("/")[-1]
                code = local_name.removesuffix(".csv")
                if selected_set and code not in selected_set:
                    continue
                download_file(SETTINGS.minio_bucket_checko, key, checko_raw_dir / local_name)
                downloaded += 1
            if downloaded == 0:
                raise ValueError("No checko CSV files matched selected OKVED codes")
            _append_job_log(
                job_uuid,
                level="INFO",
                message="Downloaded checko snapshots for stage_02",
                payload={"downloaded_files": downloaded, "selected_okved_count": len(selected_okved)},
            )
            _emit(job_id, JobStatus.RUNNING, "stage_02", 15, "Preparing active sites input")
            _run_stage(job_uuid, "stage_02_prepare_active_sites_input.py", env=env)

        total_sites = _count_stage03_total_sites(base)
        if total_sites:
            update_job_state(
                job_uuid,
                metadata_patch={
                    "runtime_progress": {
                        "processed_sites": 0,
                        "total_sites": total_sites,
                        "eta_seconds": None,
                        "current_site_url": None,
                        "updated_at": datetime.now(timezone.utc).isoformat(),
                    }
                },
            )

        _emit(job_id, JobStatus.RUNNING, "stage_03", 30, "OpenAI enrichment in progress")
        processed_urls: set[str] = set()
        stage03_started_at = time.time()
        current_site_url: str | None = None

        def _on_stage03_line(chunk: str) -> None:
            nonlocal current_site_url
            processing_match = STAGE03_PROCESSING_RE.search(chunk)
            if processing_match:
                current_site_url = processing_match.group(1).strip()
            success_match = STAGE03_SUCCESS_RE.search(chunk)
            if not success_match:
                return
            site_url = success_match.group(1).strip()
            if not site_url:
                return
            processed_urls.add(site_url)
            if total_sites:
                _update_stage03_runtime_progress(
                    job_uuid,
                    processed_sites=len(processed_urls),
                    total_sites=total_sites,
                    current_site_url=current_site_url,
                    stage_started_at=stage03_started_at,
                )

        _run_stage(job_uuid, "stage_03_enrich_sites_with_openai.py", env=env, on_line=_on_stage03_line)
        if total_sites:
            _update_stage03_runtime_progress(
                job_uuid,
                # Stage 03 is complete here, so runtime progress should be finalized to 100%
                # even if some sites were skipped/failed internally.
                processed_sites=total_sites,
                total_sites=total_sites,
                current_site_url=None,
                stage_started_at=stage03_started_at,
            )

        _emit(job_id, JobStatus.RUNNING, "stage_04", 45, "Deduplicate JSON")
        _run_stage(job_uuid, "stage_04_deduplicate_openai_json.py", env=env)

        _emit(job_id, JobStatus.RUNNING, "stage_05", 60, "Merge into Excel")
        _run_stage(job_uuid, "stage_05_merge_openai_into_excel.py", env=env)

        _emit(job_id, JobStatus.RUNNING, "stage_06", 72, "Deduplicate contacts")
        _run_stage(job_uuid, "stage_06_deduplicate_contacts.py", env=env)

        _emit(job_id, JobStatus.RUNNING, "stage_07", 84, "Normalize employees")
        _run_stage(job_uuid, "stage_07_normalize_employee_names.py", env=env)

        _emit(job_id, JobStatus.RUNNING, "stage_08", 95, "Build final sheets")
        _run_stage(job_uuid, "stage_08_add_clients_and_partners_sheets.py", env=env)

        result_file = base / "final_pipeline_output.xlsx"
        if not result_file.exists():
            result_file = base / "enriched_processed.xlsx"
        ensure_bucket(SETTINGS.minio_bucket_jobs)
        result_key = f"jobs/{job_id}/final_pipeline_output.xlsx"
        upload_file(result_file, SETTINGS.minio_bucket_jobs, result_key)

        update_job_state(
            job_uuid,
            status=JobStatus.COMPLETED,
            progress=100,
            result_bucket=SETTINGS.minio_bucket_jobs,
            result_key=result_key,
            metadata_patch=(
                {
                    "runtime_progress": {
                        "processed_sites": total_sites,
                        "total_sites": total_sites,
                        "eta_seconds": 0,
                        "current_site_url": None,
                        "updated_at": datetime.now(timezone.utc).isoformat(),
                    }
                }
                if total_sites
                else None
            ),
        )
        _emit(
            job_id,
            JobStatus.COMPLETED,
            "completed",
            100,
            "Pipeline completed",
            {"result_bucket": SETTINGS.minio_bucket_jobs, "result_key": result_key},
        )
        _append_job_log(
            job_uuid,
            level="INFO",
            message="Pipeline completed",
            payload={"result_bucket": SETTINGS.minio_bucket_jobs, "result_key": result_key},
        )
        jobs_processed_total.labels(worker="pipeline_worker", status=JobStatus.COMPLETED.value).inc()
    except JobCancelledError as exc:
        update_job_state(job_uuid, status=JobStatus.FAILED, error_message=str(exc))
        _emit(job_id, JobStatus.FAILED, "cancelled", 100, str(exc), {"cancel_requested": True})
        _append_job_log(job_uuid, level="WARNING", message=str(exc), payload={"cancel_requested": True})
        jobs_processed_total.labels(worker="pipeline_worker", status=JobStatus.FAILED.value).inc()
    except Exception as exc:
        update_job_state(job_uuid, status=JobStatus.FAILED, error_message=str(exc))
        _emit(job_id, JobStatus.FAILED, "failed", 100, f"Pipeline failed: {exc}")
        _append_job_log(job_uuid, level="ERROR", message=f"Pipeline failed: {exc}")
        jobs_processed_total.labels(worker="pipeline_worker", status=JobStatus.FAILED.value).inc()
        raise
    finally:
        HEARTBEAT.set_idle("Waiting for pipeline jobs")


HEARTBEAT = WorkerHeartbeatPublisher("pipeline_worker")


def run() -> None:
    start_http_server(int(os.getenv("METRICS_PORT", "9101")))
    ensure_schema()
    log.info("Pipeline worker started")
    HEARTBEAT.start()
    HEARTBEAT.set_idle("Waiting for pipeline jobs")
    consume_json(SETTINGS.queue_pipeline, _process, prefetch_count=1)

